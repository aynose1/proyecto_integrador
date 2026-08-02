import os
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, current_app, flash, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.api_client import APIError, api_request
from app.decorators import admin_required

bp = Blueprint("reportes", __name__, url_prefix="/reportes")

# Colores de marca (los mismos que usa toda la web: header oscuro,
# barras de nivel/peso, badges) -- se usan en ambos exportables para que
# no se sientan como una pieza aparte hecha con los defaults de la
# librería (verde genérico de antes).
COLOR_TEAL_900 = "0B3B38"
COLOR_TEAL_700 = "127A6C"
COLOR_AQUA_500 = "17B8A6"
COLOR_AQUA_100 = "E3F8F5"
COLOR_INK_900 = "13221F"


def _ruta_logo() -> str | None:
    ruta = os.path.join(current_app.static_folder, "img", "pi_fondo.png")
    return ruta if os.path.exists(ruta) else None

# 'Estado de la ruta' no es un campo guardado en la BD (solo existe
# 'completada', calculado por la API, y el estado de cada contenedor
# dentro de la ruta: pendiente/recolectado). Se deriva aquí para poder
# filtrar y mostrarlo sin tocar el backend.
ESTADOS_RUTA_LABELS = {
    "programada": "Programada",
    "en_proceso": "En proceso",
    "completada": "Completada",
}


def _parse_fecha(valor, default):
    if not valor:
        return default
    try:
        return date.fromisoformat(valor)
    except ValueError:
        return default


def _parse_hora(valor):
    if not valor:
        return None
    for formato in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(valor, formato)
        except ValueError:
            continue
    return None


def _estado_ruta(ruta: dict) -> str:
    if ruta.get("completada"):
        return "completada"
    recolectados = sum(1 for d in ruta.get("detalles", []) if d["estado"]["estado"] == "recolectado")
    return "en_proceso" if recolectados > 0 else "programada"


def _obtener_filtros():
    hoy = date.today()
    fecha_desde = _parse_fecha(request.args.get("fecha_desde"), hoy - timedelta(days=30))
    fecha_hasta = _parse_fecha(request.args.get("fecha_hasta"), hoy)
    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde

    id_sector = request.args.get("id_sector", type=int) or None
    id_recolector = request.args.get("id_recolector", type=int) or None
    estado = request.args.get("estado") or None
    if estado not in ESTADOS_RUTA_LABELS:
        estado = None

    return fecha_desde, fecha_hasta, id_sector, id_recolector, estado


def _cargar_datos_base():
    """Recolectores (para el filtro), sectores (para el filtro) y el mapa
    id->contenedor (para resolver a qué sector pertenece cada contenedor
    dentro de una ruta, dato que DetalleRuta no trae por sí solo)."""
    usuarios = api_request("get", "/usuarios")
    recolectores = [u for u in usuarios if u.get("tipo_usuario", {}).get("tipo") == "recolector"]
    sectores = api_request("get", "/sectores")
    contenedores = api_request("get", "/contenedores")
    contenedores_por_id = {c["id"]: c for c in contenedores}
    usuarios_por_id = {u["id"]: u for u in usuarios}
    return recolectores, sectores, contenedores_por_id, usuarios_por_id


def _rutas_en_rango(fecha_desde: date, fecha_hasta: date) -> list[dict]:
    """
    Una sola llamada: GET /rutas ahora acepta fecha_desde/fecha_hasta y
    devuelve el detalle completo de cada ruta (ver backend). Antes esto
    recorría día por día y pedía /rutas/{id} por cada ruta encontrada
    (N+1) — con 30 días de rango eso eran fácilmente decenas de
    solicitudes secuenciales.
    """
    return api_request(
        "get", "/rutas",
        params={
            "fecha_desde": fecha_desde.isoformat(),
            "fecha_hasta": fecha_hasta.isoformat(),
            "limit": 1000,
        },
    )


def _filtrar_rutas(rutas, id_recolector, estado, id_sector, contenedores_por_id):
    filtradas = []
    for ruta in rutas:
        if id_recolector and ruta["id_usuario"] != id_recolector:
            continue
        if estado and _estado_ruta(ruta) != estado:
            continue
        if id_sector:
            sectores_ruta = {
                contenedores_por_id[d["contenedor"]["id"]]["sector"]["id"]
                for d in ruta.get("detalles", [])
                if d["contenedor"]["id"] in contenedores_por_id
            }
            if id_sector not in sectores_ruta:
                continue
        filtradas.append(ruta)
    return filtradas


def _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
    """
    Una ruta puede incluir contenedores de más de un sector (Ruta no
    tiene 'un' sector propio en el modelo). A petición explícita, en vez
    de agrupar todo bajo 'Varios', se genera un subtotal por cada sector
    presente: eso produce una fila por sector en la tabla de detalle,
    repitiendo ruta/fecha/recolector/horario/estado en cada una.

    peso_recolectado: suma de peso_actual (la lectura más reciente que
    tiene cada contenedor ahora mismo) de los contenedores CUYO DETALLE
    quedó en 'recolectado' -- mismo criterio que ya usa 'recolectados'
    (contar), solo que sumando Kg en vez de unidades. No es el peso
    exacto que había en el momento preciso de la recolección (eso no se
    guarda en ningún lado), es una aproximación consistente con cómo ya
    funciona el resto de este reporte.
    """
    grupos = {}
    orden = []
    for d in ruta.get("detalles", []):
        cont = contenedores_por_id.get(d["contenedor"]["id"])
        if cont:
            sector_id = cont["sector"]["id"]
            sector_nombre = f'{cont["sector"]["zona"]["nombre"]} / {cont["sector"]["nombre"]}'
        else:
            sector_id = None
            sector_nombre = "Sin sector"

        if id_sector_filtro and sector_id != id_sector_filtro:
            continue

        if sector_id not in grupos:
            grupos[sector_id] = {"nombre": sector_nombre, "asignados": 0, "recolectados": 0, "peso_recolectado": 0.0}
            orden.append(sector_id)

        grupos[sector_id]["asignados"] += 1
        if d["estado"]["estado"] == "recolectado":
            grupos[sector_id]["recolectados"] += 1
            if cont:
                grupos[sector_id]["peso_recolectado"] += float(d["contenedor"].get("peso_actual") or 0)

    return [grupos[sid] for sid in orden]


def _construir_filas(rutas, contenedores_por_id, usuarios_por_id, id_sector_filtro):
    filas = []
    for ruta in rutas:
        recolector = usuarios_por_id.get(ruta["id_usuario"])
        recolector_nombre = f"{recolector['nombre']} {recolector['apellido_paterno']}" if recolector else "N/D"
        estado_ruta = ESTADOS_RUTA_LABELS[_estado_ruta(ruta)]
        fecha_ruta = date.fromisoformat(ruta["fecha"])

        for grupo in _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
            filas.append(
                {
                    "ruta": ruta["nombre"],
                    "fecha": fecha_ruta,
                    "recolector": recolector_nombre,
                    "sector": grupo["nombre"],
                    "asignados": grupo["asignados"],
                    "recolectados": grupo["recolectados"],
                    "peso_recolectado": round(grupo["peso_recolectado"], 2),
                    "hora_inicio": ruta.get("hora_inicio"),
                    "hora_fin": ruta.get("hora_fin"),
                    "estado": estado_ruta,
                }
            )
    return filas


def _resumen_general(rutas, id_sector_filtro, contenedores_por_id):
    total_rutas = len(rutas)
    completadas = sum(1 for r in rutas if r.get("completada"))

    asignados = 0
    recolectados = 0
    peso_total = 0.0
    for ruta in rutas:
        for grupo in _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
            asignados += grupo["asignados"]
            recolectados += grupo["recolectados"]
            peso_total += grupo["peso_recolectado"]

    porcentaje = round((recolectados / asignados * 100), 1) if asignados else 0.0

    return {
        "total_rutas": total_rutas,
        "rutas_completadas": completadas,
        "contenedores_asignados": asignados,
        "contenedores_recolectados": recolectados,
        "porcentaje_cumplimiento": porcentaje,
        "peso_total_recolectado": round(peso_total, 2),
    }


def _resumen_final(rutas, id_sector_filtro, contenedores_por_id):
    """
    "Resumen" (antes "Resumen final"): reemplaza por completo las
    métricas de tiempo/tamaño de ruta -- ahora es puramente sobre peso
    recolectado en el periodo:
      - peso_total_kg: la misma suma que ya calcula _resumen_general,
        recalculada aquí para que esta función sea autosuficiente.
      - peso_promedio_por_ruta_kg: peso_total entre RUTAS COMPLETADAS
        (no entre todas -- una ruta sin terminar no representa un
        "recorrido típico" completo).
      - sector_mayor_peso: qué sector aportó más peso en el periodo,
        agregando todos los grupos de todas las rutas por sector.
      - peso_promedio_por_contenedor_kg: peso_total entre el número de
        contenedores efectivamente recolectados (indica qué tan lleno
        llega en promedio cada contenedor al momento de recogerlo).
    """
    peso_total = 0.0
    completadas = 0
    total_recolectados = 0
    peso_por_sector: dict[str, float] = {}

    for ruta in rutas:
        if ruta.get("completada"):
            completadas += 1
        for grupo in _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
            peso_total += grupo["peso_recolectado"]
            total_recolectados += grupo["recolectados"]
            if grupo["peso_recolectado"] > 0:
                peso_por_sector[grupo["nombre"]] = peso_por_sector.get(grupo["nombre"], 0.0) + grupo["peso_recolectado"]

    peso_promedio_ruta = round(peso_total / completadas, 2) if completadas else None
    peso_promedio_contenedor = round(peso_total / total_recolectados, 2) if total_recolectados else None
    sector_mayor = max(peso_por_sector.items(), key=lambda x: x[1], default=(None, 0.0))

    return {
        "peso_total_kg": round(peso_total, 2),
        "peso_promedio_por_ruta_kg": peso_promedio_ruta,
        "sector_mayor_peso_nombre": sector_mayor[0],
        "sector_mayor_peso_kg": round(sector_mayor[1], 2) if sector_mayor[0] else None,
        "peso_promedio_por_contenedor_kg": peso_promedio_contenedor,
    }


def _col_widths(ancho_disponible, pesos):
    """Reparte ancho_disponible entre columnas según pesos relativos, para
    que la tabla ocupe todo el ancho útil de la página en vez de quedar
    encogida al tamaño de su contenido."""
    total_pesos = sum(pesos)
    return [ancho_disponible * peso / total_pesos for peso in pesos]


def _contexto_reporte():
    fecha_desde, fecha_hasta, id_sector, id_recolector, estado = _obtener_filtros()
    recolectores, sectores, contenedores_por_id, usuarios_por_id = _cargar_datos_base()

    rutas_rango = _rutas_en_rango(fecha_desde, fecha_hasta)
    rutas = _filtrar_rutas(rutas_rango, id_recolector, estado, id_sector, contenedores_por_id)

    return {
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "id_sector": id_sector,
        "id_recolector": id_recolector,
        "estado": estado,
        "recolectores": recolectores,
        "sectores": sectores,
        "filas": _construir_filas(rutas, contenedores_por_id, usuarios_por_id, id_sector),
        "resumen": _resumen_general(rutas, id_sector, contenedores_por_id),
        "resumen_final": _resumen_final(rutas, id_sector, contenedores_por_id),
        "generado_en": datetime.now(),
    }


@bp.route("")
@admin_required
def index():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        contexto = None

    return render_template(
        "reportes/index.html",
        contexto=contexto,
        institucion=current_app.config["INSTITUTION_NAME"],
        estados_ruta=ESTADOS_RUTA_LABELS,
    )


@bp.route("/exportar/excel")
@admin_required
def exportar_excel():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        return redirect(url_for("reportes.index"))

    institucion = current_app.config["INSTITUTION_NAME"]
    r = contexto["resumen"]
    rf = contexto["resumen_final"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte General"

    # ---------- Estilos reutilizables ----------
    relleno_teal = PatternFill("solid", fgColor=COLOR_TEAL_900)
    relleno_aqua_claro = PatternFill("solid", fgColor=COLOR_AQUA_100)
    fuente_titulo = Font(bold=True, size=16, color=COLOR_TEAL_900)
    fuente_subtitulo = Font(size=10, color="4C6360")
    fuente_encabezado_seccion = Font(bold=True, size=11, color="FFFFFF")
    fuente_encabezado_tabla = Font(bold=True, size=10, color="FFFFFF")
    fuente_etiqueta = Font(bold=True, size=10, color=COLOR_INK_900)
    borde_delgado = Border(*[Side(style="thin", color="D9E4E2")] * 4)

    def _fila_seccion(texto, ultima_columna):
        """Una fila completa rellena de teal con el texto en blanco -- usada
        para los 3 encabezados de sección (Resumen general, Detalle, Resumen)."""
        fila = ws.max_row + 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_columna)
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.font = fuente_encabezado_seccion
        celda.fill = relleno_teal
        celda.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[fila].height = 20
        return fila

    # ---------- Encabezado con logo ----------
    ruta_logo = _ruta_logo()
    if ruta_logo:
        img = XLImage(ruta_logo)
        img.width = 90
        img.height = 70
        ws.add_image(img, "A1")

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 16

    ws["C1"] = institucion
    ws["C1"].font = fuente_titulo
    ws["C2"] = "Sistema de Gestión de Recolección de Residuos Sólidos"
    ws["C2"].font = fuente_subtitulo
    ws["C3"] = f"Reporte General de Recolección de Rutas — Generado: {contexto['generado_en'].strftime('%d/%m/%Y %H:%M')}"
    ws["C3"].font = fuente_subtitulo
    ws["C4"] = (
        f"Periodo consultado: {contexto['fecha_desde'].strftime('%d/%m/%Y')} — "
        f"{contexto['fecha_hasta'].strftime('%d/%m/%Y')}"
    )
    ws["C4"].font = fuente_subtitulo
    ws.append([])

    # ---------- Resumen general ----------
    _fila_seccion("Resumen general", 2)
    resumen_general_filas = [
        ("Total de rutas programadas", r["total_rutas"], None),
        ("Total de rutas completadas", r["rutas_completadas"], None),
        ("Total de contenedores asignados", r["contenedores_asignados"], None),
        ("Total de contenedores recolectados", r["contenedores_recolectados"], None),
        ("Porcentaje de cumplimiento", r["porcentaje_cumplimiento"] / 100, "0.0%"),
        ("Peso total recolectado (Kg)", r["peso_total_recolectado"], "0.00"),
    ]
    for etiqueta, valor, formato in resumen_general_filas:
        fila = ws.max_row + 1
        c_etq = ws.cell(row=fila, column=1, value=etiqueta)
        c_val = ws.cell(row=fila, column=2, value=valor)
        c_etq.font = fuente_etiqueta
        c_etq.border = borde_delgado
        c_val.border = borde_delgado
        if formato:
            c_val.number_format = formato
    ws.append([])

    # ---------- Detalle de rutas ----------
    fila_seccion_detalle = _fila_seccion("Detalle de rutas", 10)
    fila_encabezado = ws.max_row + 1
    encabezados = [
        "Ruta", "Fecha", "Recolector", "Sector", "Contenedores asignados",
        "Contenedores recolectados", "Peso recolectado (Kg)", "Hora inicio", "Hora fin", "Estado",
    ]
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=col, value=texto)
        celda.font = fuente_encabezado_tabla
        celda.fill = PatternFill("solid", fgColor=COLOR_TEAL_700)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        celda.border = borde_delgado

    for i, fila_datos in enumerate(contexto["filas"]):
        fila = ws.max_row + 1
        valores = [
            fila_datos["ruta"], fila_datos["fecha"], fila_datos["recolector"], fila_datos["sector"],
            fila_datos["asignados"], fila_datos["recolectados"], fila_datos["peso_recolectado"],
            fila_datos["hora_inicio"] or "-", fila_datos["hora_fin"] or "-", fila_datos["estado"],
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.border = borde_delgado
            if i % 2 == 1:
                celda.fill = relleno_aqua_claro
        ws.cell(row=fila, column=2).number_format = "dd/mm/yyyy"
        ws.cell(row=fila, column=7).number_format = "0.00"

    if not contexto["filas"]:
        ws.append(["Sin rutas para los filtros seleccionados"])
    ws.append([])

    # Encabezado de la tabla de detalle siempre visible al hacer scroll,
    # y filtro rápido en cada columna.
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)
    ws.auto_filter.ref = f"A{fila_encabezado}:J{fila_encabezado}"

    # ---------- Resumen (peso) ----------
    _fila_seccion("Resumen", 2)
    resumen_final_filas = [
        ("Peso total recolectado (Kg)", rf["peso_total_kg"], "0.00"),
        (
            "Peso promedio por ruta completada (Kg)",
            rf["peso_promedio_por_ruta_kg"] if rf["peso_promedio_por_ruta_kg"] is not None else "No disponible",
            "0.00" if rf["peso_promedio_por_ruta_kg"] is not None else None,
        ),
        (
            "Sector con mayor peso recolectado",
            f"{rf['sector_mayor_peso_nombre']} ({rf['sector_mayor_peso_kg']} Kg)" if rf["sector_mayor_peso_nombre"] else "N/D",
            None,
        ),
        (
            "Peso promedio por contenedor recolectado (Kg)",
            rf["peso_promedio_por_contenedor_kg"] if rf["peso_promedio_por_contenedor_kg"] is not None else "No disponible",
            "0.00" if rf["peso_promedio_por_contenedor_kg"] is not None else None,
        ),
    ]
    for etiqueta, valor, formato in resumen_final_filas:
        fila = ws.max_row + 1
        c_etq = ws.cell(row=fila, column=1, value=etiqueta)
        c_val = ws.cell(row=fila, column=2, value=valor)
        c_etq.font = fuente_etiqueta
        c_etq.border = borde_delgado
        c_val.border = borde_delgado
        if formato:
            c_val.number_format = formato

    anchos = (30, 13, 22, 26, 14, 16, 15, 12, 12, 14)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_general_recoleccion_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class _NumberedCanvas(canvas.Canvas):
    """
    Necesario para poder poner "Página X de Y" -- reportlab no sabe
    cuántas páginas habrá hasta terminar de construir el documento
    completo, así que se guarda el estado de cada página según se van
    generando, y hasta el final (save()) se vuelve a pasar por todas
    para escribir el número total ya conocido.
    """
    def __init__(self, *args, **kwargs):
        canvas.Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total_paginas = len(self._saved_page_states)
        for estado in self._saved_page_states:
            self.__dict__.update(estado)
            self._dibujar_numero_pagina(total_paginas)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def _dibujar_numero_pagina(self, total_paginas):
        ancho_pagina, _ = landscape(letter)
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor(f"#{COLOR_INK_900}"))
        self.drawRightString(
            ancho_pagina - 0.5 * inch, 0.28 * inch, f"Página {self._pageNumber} de {total_paginas}"
        )


def _dibujar_marco_pagina(canvas_obj, doc_obj, institucion, ruta_logo, generado_texto):
    """
    Se llama automáticamente por reportlab en CADA página (onFirstPage/
    onLaterPages) -- dibuja la franja superior de color de marca con el
    logo, y el pie de página con la fecha de generación. El número de
    página ("de Y") lo pone _NumberedCanvas por separado, porque este
    callback no sabe todavía cuántas páginas habrá en total.
    """
    ancho_pagina, alto_pagina = landscape(letter)
    alto_franja = 0.85 * inch

    canvas_obj.saveState()

    # Franja superior
    canvas_obj.setFillColor(colors.HexColor(f"#{COLOR_TEAL_900}"))
    canvas_obj.rect(0, alto_pagina - alto_franja, ancho_pagina, alto_franja, fill=1, stroke=0)

    if ruta_logo:
        try:
            canvas_obj.drawImage(
                ruta_logo, 0.35 * inch, alto_pagina - alto_franja + 0.1 * inch,
                width=0.65 * inch, height=0.65 * inch,
                preserveAspectRatio=True, mask="auto",
            )
        except Exception:
            pass  # si el logo no se puede leer, el reporte se genera igual, solo sin imagen

    texto_x = 1.2 * inch
    canvas_obj.setFillColor(colors.white)
    canvas_obj.setFont("Helvetica-Bold", 14)
    canvas_obj.drawString(texto_x, alto_pagina - 0.38 * inch, institucion)
    canvas_obj.setFont("Helvetica", 9)
    canvas_obj.setFillColor(colors.HexColor(f"#{COLOR_AQUA_500}"))
    canvas_obj.drawString(texto_x, alto_pagina - 0.58 * inch, "Reporte General de Recolección de Rutas")

    # Pie de página
    canvas_obj.setFillColor(colors.HexColor("#7D928E"))
    canvas_obj.setFont("Helvetica", 8)
    canvas_obj.drawString(0.5 * inch, 0.28 * inch, generado_texto)
    canvas_obj.setStrokeColor(colors.HexColor("#D9E4E2"))
    canvas_obj.line(0.5 * inch, 0.45 * inch, ancho_pagina - 0.5 * inch, 0.45 * inch)

    canvas_obj.restoreState()


@bp.route("/exportar/pdf")
@admin_required
def exportar_pdf():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        return redirect(url_for("reportes.index"))

    institucion = current_app.config["INSTITUTION_NAME"]
    r = contexto["resumen"]
    rf = contexto["resumen_final"]
    ruta_logo = _ruta_logo()
    generado_texto = (
        f"Generado: {contexto['generado_en'].strftime('%d/%m/%Y %H:%M')}  ·  "
        f"Periodo: {contexto['fecha_desde'].strftime('%d/%m/%Y')} — {contexto['fecha_hasta'].strftime('%d/%m/%Y')}"
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        title="Reporte General de Recolección de Rutas",
        topMargin=1.05 * inch, bottomMargin=0.6 * inch,
        leftMargin=0.5 * inch, rightMargin=0.5 * inch,
    )
    ancho_disponible = doc.width  # ancho de página menos leftMargin/rightMargin
    styles = getSampleStyleSheet()
    estilo_seccion = ParagraphStyle(
        "Seccion", parent=styles["Heading3"], textColor=colors.HexColor(f"#{COLOR_TEAL_900}"), spaceAfter=6,
    )
    elementos = []

    color_teal = colors.HexColor(f"#{COLOR_TEAL_900}")
    color_teal_claro = colors.HexColor(f"#{COLOR_TEAL_700}")
    color_aqua_claro = colors.HexColor(f"#{COLOR_AQUA_100}")
    color_borde = colors.HexColor("#D9E4E2")

    elementos.append(Paragraph("Resumen general", estilo_seccion))
    resumen_data = [
        ["Total de rutas programadas", str(r["total_rutas"])],
        ["Total de rutas completadas", str(r["rutas_completadas"])],
        ["Total de contenedores asignados", str(r["contenedores_asignados"])],
        ["Total de contenedores recolectados", str(r["contenedores_recolectados"])],
        ["Porcentaje de cumplimiento", f"{r['porcentaje_cumplimiento']}%"],
        ["Peso total recolectado", f"{r['peso_total_recolectado']} Kg"],
    ]
    tabla_resumen = Table(resumen_data, colWidths=_col_widths(ancho_disponible, [3, 1]))
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), color_aqua_claro),
        ("TEXTCOLOR", (0, 0), (0, -1), color_teal),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, color_borde),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(tabla_resumen)
    elementos.append(Spacer(1, 0.3 * inch))

    elementos.append(Paragraph("Detalle de rutas", estilo_seccion))
    encabezados = ["Ruta", "Fecha", "Recolector", "Sector", "Asignados", "Recolectados", "Peso (Kg)", "H. inicio", "H. fin", "Estado"]
    filas_tabla = [encabezados]
    for fila in contexto["filas"]:
        filas_tabla.append([
            fila["ruta"], fila["fecha"].strftime("%d/%m/%Y"), fila["recolector"], fila["sector"],
            str(fila["asignados"]), str(fila["recolectados"]), str(fila["peso_recolectado"]),
            fila["hora_inicio"] or "-", fila["hora_fin"] or "-", fila["estado"],
        ])
    if len(filas_tabla) == 1:
        filas_tabla.append(["Sin rutas para los filtros seleccionados"] + [""] * 9)

    pesos_detalle = [15, 8, 14, 15, 8, 9, 9, 8, 8, 9]  # Ruta, Fecha, Recolector, Sector, Asig., Recol., Peso, H.ini, H.fin, Estado
    tabla_detalle = Table(filas_tabla, colWidths=_col_widths(ancho_disponible, pesos_detalle), repeatRows=1)
    tabla_detalle.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), color_teal),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.5, color_borde),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, color_aqua_claro]),
    ]))
    elementos.append(tabla_detalle)
    elementos.append(Spacer(1, 0.3 * inch))

    elementos.append(Paragraph("Resumen", estilo_seccion))
    resumen_final_data = [
        ["Peso total recolectado", f"{rf['peso_total_kg']} Kg"],
        [
            "Peso promedio por ruta completada",
            f"{rf['peso_promedio_por_ruta_kg']} Kg" if rf["peso_promedio_por_ruta_kg"] is not None else "No disponible",
        ],
        [
            "Sector con mayor peso recolectado",
            f"{rf['sector_mayor_peso_nombre']} ({rf['sector_mayor_peso_kg']} Kg)" if rf["sector_mayor_peso_nombre"] else "N/D",
        ],
        [
            "Peso promedio por contenedor recolectado",
            f"{rf['peso_promedio_por_contenedor_kg']} Kg" if rf["peso_promedio_por_contenedor_kg"] is not None else "No disponible",
        ],
    ]
    tabla_final = Table(resumen_final_data, colWidths=_col_widths(ancho_disponible, [3, 2]))
    tabla_final.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), color_aqua_claro),
        ("TEXTCOLOR", (0, 0), (0, -1), color_teal),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, color_borde),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(tabla_final)

    def _marco(canvas_obj, doc_obj):
        _dibujar_marco_pagina(canvas_obj, doc_obj, institucion, ruta_logo, generado_texto)

    doc.build(elementos, onFirstPage=_marco, onLaterPages=_marco, canvasmaker=_NumberedCanvas)
    buffer.seek(0)

    filename = f"reporte_general_recoleccion_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/pdf")

