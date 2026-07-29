from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, current_app, flash, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.api_client import APIError, api_request
from app.decorators import admin_required

bp = Blueprint("reportes", __name__, url_prefix="/reportes")

# 'Estado de la ruta' no es un campo guardado en la BD (solo existe
# 'completada', calculado por la API, y el estado de cada contenedor
# dentro de la ruta: pendiente/recolectado). Se deriva aquí para poder
# filtrar y mostrarlo sin tocar el backend.
ESTADOS_RUTA_LABELS = {
    "programada": "Programada",
    "en_proceso": "En proceso",
    "completada": "Completada",
}


def _parse_fecha(valor, default):
    if not valor:
        return default
    try:
        return date.fromisoformat(valor)
    except ValueError:
        return default


def _parse_hora(valor):
    if not valor:
        return None
    for formato in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(valor, formato)
        except ValueError:
            continue
    return None


def _estado_ruta(ruta: dict) -> str:
    if ruta.get("completada"):
        return "completada"
    recolectados = sum(1 for d in ruta.get("detalles", []) if d["estado"]["estado"] == "recolectado")
    return "en_proceso" if recolectados > 0 else "programada"


def _obtener_filtros():
    hoy = date.today()
    fecha_desde = _parse_fecha(request.args.get("fecha_desde"), hoy - timedelta(days=30))
    fecha_hasta = _parse_fecha(request.args.get("fecha_hasta"), hoy)
    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde

    id_sector = request.args.get("id_sector", type=int) or None
    id_recolector = request.args.get("id_recolector", type=int) or None
    estado = request.args.get("estado") or None
    if estado not in ESTADOS_RUTA_LABELS:
        estado = None

    return fecha_desde, fecha_hasta, id_sector, id_recolector, estado


def _cargar_datos_base():
    """Recolectores (para el filtro), sectores (para el filtro) y el mapa
    id->contenedor (para resolver a qué sector pertenece cada contenedor
    dentro de una ruta, dato que DetalleRuta no trae por sí solo)."""
    usuarios = api_request("get", "/usuarios")
    recolectores = [u for u in usuarios if u.get("tipo_usuario", {}).get("tipo") == "recolector"]
    sectores = api_request("get", "/sectores")
    contenedores = api_request("get", "/contenedores")
    contenedores_por_id = {c["id"]: c for c in contenedores}
    usuarios_por_id = {u["id"]: u for u in usuarios}
    return recolectores, sectores, contenedores_por_id, usuarios_por_id


def _rutas_en_rango(fecha_desde: date, fecha_hasta: date) -> list[dict]:
    """
    Una sola llamada: GET /rutas ahora acepta fecha_desde/fecha_hasta y
    devuelve el detalle completo de cada ruta (ver backend). Antes esto
    recorría día por día y pedía /rutas/{id} por cada ruta encontrada
    (N+1) — con 30 días de rango eso eran fácilmente decenas de
    solicitudes secuenciales.
    """
    return api_request(
        "get", "/rutas",
        params={
            "fecha_desde": fecha_desde.isoformat(),
            "fecha_hasta": fecha_hasta.isoformat(),
            "limit": 1000,
        },
    )


def _filtrar_rutas(rutas, id_recolector, estado, id_sector, contenedores_por_id):
    filtradas = []
    for ruta in rutas:
        if id_recolector and ruta["id_usuario"] != id_recolector:
            continue
        if estado and _estado_ruta(ruta) != estado:
            continue
        if id_sector:
            sectores_ruta = {
                contenedores_por_id[d["contenedor"]["id"]]["sector"]["id"]
                for d in ruta.get("detalles", [])
                if d["contenedor"]["id"] in contenedores_por_id
            }
            if id_sector not in sectores_ruta:
                continue
        filtradas.append(ruta)
    return filtradas


def _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
    """
    Una ruta puede incluir contenedores de más de un sector (Ruta no
    tiene 'un' sector propio en el modelo). A petición explícita, en vez
    de agrupar todo bajo 'Varios', se genera un subtotal por cada sector
    presente: eso produce una fila por sector en la tabla de detalle,
    repitiendo ruta/fecha/recolector/horario/estado en cada una.
    """
    grupos = {}
    orden = []
    for d in ruta.get("detalles", []):
        cont = contenedores_por_id.get(d["contenedor"]["id"])
        if cont:
            sector_id = cont["sector"]["id"]
            sector_nombre = f'{cont["sector"]["zona"]["nombre"]} / {cont["sector"]["nombre"]}'
        else:
            sector_id = None
            sector_nombre = "Sin sector"

        if id_sector_filtro and sector_id != id_sector_filtro:
            continue

        if sector_id not in grupos:
            grupos[sector_id] = {"nombre": sector_nombre, "asignados": 0, "recolectados": 0}
            orden.append(sector_id)

        grupos[sector_id]["asignados"] += 1
        if d["estado"]["estado"] == "recolectado":
            grupos[sector_id]["recolectados"] += 1

    return [grupos[sid] for sid in orden]


def _construir_filas(rutas, contenedores_por_id, usuarios_por_id, id_sector_filtro):
    filas = []
    for ruta in rutas:
        recolector = usuarios_por_id.get(ruta["id_usuario"])
        recolector_nombre = f"{recolector['nombre']} {recolector['apellido_paterno']}" if recolector else "N/D"
        estado_ruta = ESTADOS_RUTA_LABELS[_estado_ruta(ruta)]
        fecha_ruta = date.fromisoformat(ruta["fecha"])

        for grupo in _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
            filas.append(
                {
                    "ruta": ruta["nombre"],
                    "fecha": fecha_ruta,
                    "recolector": recolector_nombre,
                    "sector": grupo["nombre"],
                    "asignados": grupo["asignados"],
                    "recolectados": grupo["recolectados"],
                    "hora_inicio": ruta.get("hora_inicio"),
                    "hora_fin": ruta.get("hora_fin"),
                    "estado": estado_ruta,
                }
            )
    return filas


def _resumen_general(rutas, id_sector_filtro, contenedores_por_id):
    total_rutas = len(rutas)
    completadas = sum(1 for r in rutas if r.get("completada"))

    asignados = 0
    recolectados = 0
    for ruta in rutas:
        for grupo in _agrupar_por_sector(ruta, contenedores_por_id, id_sector_filtro):
            asignados += grupo["asignados"]
            recolectados += grupo["recolectados"]

    porcentaje = round((recolectados / asignados * 100), 1) if asignados else 0.0

    return {
        "total_rutas": total_rutas,
        "rutas_completadas": completadas,
        "contenedores_asignados": asignados,
        "contenedores_recolectados": recolectados,
        "porcentaje_cumplimiento": porcentaje,
    }


def _resumen_final(rutas):
    """
    hora_inicio/hora_fin son opcionales en el modelo (Ruta puede no
    tenerlas capturadas), así que el promedio solo considera las rutas
    que sí las tienen; si ninguna las tiene, se reporta 'no disponible'
    en vez de fallar.
    """
    tiempos = []
    for ruta in rutas:
        t1 = _parse_hora(ruta.get("hora_inicio"))
        t2 = _parse_hora(ruta.get("hora_fin"))
        if t1 and t2:
            minutos = (t2 - t1).total_seconds() / 60
            if minutos > 0:
                tiempos.append(minutos)

    tiempo_promedio = round(sum(tiempos) / len(tiempos)) if tiempos else None

    con_cantidad = [(r, len(r.get("detalles", []))) for r in rutas]
    ruta_mayor = max(con_cantidad, key=lambda x: x[1], default=(None, 0))
    ruta_menor = min(con_cantidad, key=lambda x: x[1], default=(None, 0))

    return {
        "tiempo_promedio_minutos": tiempo_promedio,
        "ruta_mayor": ruta_mayor[0]["nombre"] if ruta_mayor[0] else None,
        "ruta_mayor_cantidad": ruta_mayor[1],
        "ruta_menor": ruta_menor[0]["nombre"] if ruta_menor[0] else None,
        "ruta_menor_cantidad": ruta_menor[1],
    }


def _col_widths(ancho_disponible, pesos):
    """Reparte ancho_disponible entre columnas según pesos relativos, para
    que la tabla ocupe todo el ancho útil de la página en vez de quedar
    encogida al tamaño de su contenido."""
    total_pesos = sum(pesos)
    return [ancho_disponible * peso / total_pesos for peso in pesos]


def _contexto_reporte():
    fecha_desde, fecha_hasta, id_sector, id_recolector, estado = _obtener_filtros()
    recolectores, sectores, contenedores_por_id, usuarios_por_id = _cargar_datos_base()

    rutas_rango = _rutas_en_rango(fecha_desde, fecha_hasta)
    rutas = _filtrar_rutas(rutas_rango, id_recolector, estado, id_sector, contenedores_por_id)

    return {
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "id_sector": id_sector,
        "id_recolector": id_recolector,
        "estado": estado,
        "recolectores": recolectores,
        "sectores": sectores,
        "filas": _construir_filas(rutas, contenedores_por_id, usuarios_por_id, id_sector),
        "resumen": _resumen_general(rutas, id_sector, contenedores_por_id),
        "resumen_final": _resumen_final(rutas),
        "generado_en": datetime.now(),
    }


@bp.route("")
@admin_required
def index():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        contexto = None

    return render_template(
        "reportes/index.html",
        contexto=contexto,
        institucion=current_app.config["INSTITUTION_NAME"],
        estados_ruta=ESTADOS_RUTA_LABELS,
    )


@bp.route("/exportar/excel")
@admin_required
def exportar_excel():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        return redirect(url_for("reportes.index"))

    institucion = current_app.config["INSTITUTION_NAME"]
    r = contexto["resumen"]
    rf = contexto["resumen_final"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte General"

    ws.append([institucion])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Sistema de Gestión de Recolección de Residuos Sólidos"])
    ws.append(["Reporte General de Recolección de Rutas"])
    ws.append([f"Generado: {contexto['generado_en'].strftime('%d/%m/%Y %H:%M')}"])
    ws.append([
        f"Periodo consultado: {contexto['fecha_desde'].strftime('%d/%m/%Y')} — "
        f"{contexto['fecha_hasta'].strftime('%d/%m/%Y')}"
    ])
    ws.append([])

    ws.append(["Resumen general"])
    ws["A7"].font = Font(bold=True)
    ws.append(["Total de rutas programadas", r["total_rutas"]])
    ws.append(["Total de rutas completadas", r["rutas_completadas"]])
    ws.append(["Total de contenedores asignados", r["contenedores_asignados"]])
    ws.append(["Total de contenedores recolectados", r["contenedores_recolectados"]])
    ws.append(["Porcentaje de cumplimiento", f"{r['porcentaje_cumplimiento']}%"])
    ws.append([])

    fila_encabezado = ws.max_row + 1
    ws.append([
        "Ruta", "Fecha", "Recolector", "Sector", "Contenedores asignados",
        "Contenedores recolectados", "Hora inicio", "Hora fin", "Estado",
    ])
    for celda in ws[fila_encabezado]:
        celda.font = Font(bold=True)

    for fila in contexto["filas"]:
        ws.append([
            fila["ruta"], fila["fecha"], fila["recolector"], fila["sector"],
            fila["asignados"], fila["recolectados"],
            fila["hora_inicio"] or "-", fila["hora_fin"] or "-", fila["estado"],
        ])
    if not contexto["filas"]:
        ws.append(["Sin rutas para los filtros seleccionados"])
    ws.append([])

    ws.append(["Resumen final"])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    ws.append([
        "Tiempo promedio por ruta (min)",
        rf["tiempo_promedio_minutos"] if rf["tiempo_promedio_minutos"] is not None else "No disponible",
    ])
    ws.append([
        "Ruta con mayor número de contenedores",
        f"{rf['ruta_mayor']} ({rf['ruta_mayor_cantidad']})" if rf["ruta_mayor"] else "N/D",
    ])
    ws.append([
        "Ruta con menor número de contenedores",
        f"{rf['ruta_menor']} ({rf['ruta_menor_cantidad']})" if rf["ruta_menor"] else "N/D",
    ])

    for columna, ancho in zip("ABCDEFGHI", (28, 12, 22, 26, 14, 14, 12, 12, 14)):
        ws.column_dimensions[columna].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_general_recoleccion_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/exportar/pdf")
@admin_required
def exportar_pdf():
    try:
        contexto = _contexto_reporte()
    except APIError as exc:
        flash(exc.message, "danger")
        return redirect(url_for("reportes.index"))

    institucion = current_app.config["INSTITUTION_NAME"]
    r = contexto["resumen"]
    rf = contexto["resumen_final"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        title="Reporte General de Recolección de Rutas",
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
        leftMargin=0.5 * inch, rightMargin=0.5 * inch,
    )
    ancho_disponible = doc.width  # ancho de página menos leftMargin/rightMargin
    styles = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph(institucion, styles["Title"]))
    elementos.append(Paragraph("Sistema de Gestión de Recolección de Residuos Sólidos", styles["Heading2"]))
    elementos.append(Paragraph("Reporte General de Recolección de Rutas", styles["Heading3"]))
    elementos.append(Paragraph(f"Generado: {contexto['generado_en'].strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    elementos.append(Paragraph(
        f"Periodo consultado: {contexto['fecha_desde'].strftime('%d/%m/%Y')} — "
        f"{contexto['fecha_hasta'].strftime('%d/%m/%Y')}",
        styles["Normal"],
    ))
    elementos.append(Spacer(1, 0.25 * inch))

    elementos.append(Paragraph("Resumen general", styles["Heading3"]))
    resumen_data = [
        ["Total de rutas programadas", str(r["total_rutas"])],
        ["Total de rutas completadas", str(r["rutas_completadas"])],
        ["Total de contenedores asignados", str(r["contenedores_asignados"])],
        ["Total de contenedores recolectados", str(r["contenedores_recolectados"])],
        ["Porcentaje de cumplimiento", f"{r['porcentaje_cumplimiento']}%"],
    ]
    tabla_resumen = Table(resumen_data, colWidths=_col_widths(ancho_disponible, [3, 1]))
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.whitesmoke),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabla_resumen)
    elementos.append(Spacer(1, 0.3 * inch))

    elementos.append(Paragraph("Detalle de rutas", styles["Heading3"]))
    encabezados = ["Ruta", "Fecha", "Recolector", "Sector", "Asignados", "Recolectados", "H. inicio", "H. fin", "Estado"]
    filas_tabla = [encabezados]
    for fila in contexto["filas"]:
        filas_tabla.append([
            fila["ruta"], fila["fecha"].strftime("%d/%m/%Y"), fila["recolector"], fila["sector"],
            str(fila["asignados"]), str(fila["recolectados"]),
            fila["hora_inicio"] or "-", fila["hora_fin"] or "-", fila["estado"],
        ])
    if len(filas_tabla) == 1:
        filas_tabla.append(["Sin rutas para los filtros seleccionados"] + [""] * 8)

    pesos_detalle = [16, 9, 15, 16, 9, 10, 9, 9, 10]  # Ruta, Fecha, Recolector, Sector, Asig., Recol., H.ini, H.fin, Estado
    tabla_detalle = Table(filas_tabla, colWidths=_col_widths(ancho_disponible, pesos_detalle), repeatRows=1)
    tabla_detalle.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e7d32")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f4")]),
    ]))
    elementos.append(tabla_detalle)
    elementos.append(Spacer(1, 0.3 * inch))

    elementos.append(Paragraph("Resumen final", styles["Heading3"]))
    resumen_final_data = [
        [
            "Tiempo promedio por ruta",
            f"{rf['tiempo_promedio_minutos']} min" if rf["tiempo_promedio_minutos"] is not None else "No disponible",
        ],
        [
            "Ruta con mayor número de contenedores",
            f"{rf['ruta_mayor']} ({rf['ruta_mayor_cantidad']})" if rf["ruta_mayor"] else "N/D",
        ],
        [
            "Ruta con menor número de contenedores",
            f"{rf['ruta_menor']} ({rf['ruta_menor_cantidad']})" if rf["ruta_menor"] else "N/D",
        ],
    ]
    tabla_final = Table(resumen_final_data, colWidths=_col_widths(ancho_disponible, [3, 2]))
    tabla_final.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elementos.append(tabla_final)

    doc.build(elementos)
    buffer.seek(0)

    filename = f"reporte_general_recoleccion_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/pdf")
